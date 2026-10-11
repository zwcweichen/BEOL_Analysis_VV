# -*- coding: utf-8 -*-
"""建立 docs/REVIEW_LOG.xlsx（三角色自審紀錄）的初始版本。
之後這份 Excel 是「活的」：每一版自審直接在 Excel 內加列；這個檔只用來第一次建立或重建格式（會覆蓋！）。
執行：python docs/build_review_log.py   （需要 openpyxl）"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'REVIEW_LOG.xlsx')
HEAD = PatternFill('solid', fgColor='1F3B63'); HFONT = Font(bold=True, color='FFFFFF')
THIN = Side(style='thin', color='D0D5DD'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(wb, title, cols, rows, first=True):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.fill = HEAD; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v); c.alignment = WRAP; c.border = BORDER
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30
    return ws


wb = Workbook()

# ---- 1. 角色定義
roles = [
    ['A', '開發（與使用者討論）', '製程整合主管／資深管理 + 系統程式設計',
     '把需求化成最簡潔的設計：避免複雜 UI、固定分析流程不發散、確保製程分析真正需要的輸出（點燈、HL、原因、檢查、跨模組關係、bad tool）都在。',
     '每版開始：整理回饋 → 設計 → 實作 → 測試。每版結束：讀 B、C 的發現，決定「修／延後／不做」並記錄理由。'],
    ['B', '使用者模擬（工程師）', '製程整合工程師，手上只有自己的資料',
     '每版產生 1～2 份模擬資料（可取材知識庫文獻的 issue：例如 dishing 繼承、via 下 SIV、tool 偏移、無機台欄、長表、不同欄名），當作「新使用者的資料」丟進工具，目的是驗證泛用性。',
     '只做「放資料 → 看輸出」，不看原始碼。把模擬資料與埋入的故事寫進「模擬資料」工作表。'],
    ['C', '使用者體驗回饋（工程師 + 管理者）', '與 B 相同視角的工程師，再加一個管理者視角',
     '放入自己的資料後，逐項檢查：有沒有卡住（錯誤、空白、等待）、欄位抓不抓得到（比較欄、分組欄、wafer id、item 類別）、能不能更方便（少一次點擊、少一段文字）、結論是否能直接拿去跟主管講（管理者視角：一眼看懂、可追溯、不過度推論）。',
     '每個發現一列寫進「自審紀錄」，附類型與嚴重度；不直接改程式。'],
]
sheet(wb, '角色定義', [('角色', 6), ('名稱', 22), ('視角', 30), ('職責', 70), ('做法', 60)], roles)

# ---- 2. 自審紀錄（活的）
log_cols = [('版本', 8), ('日期', 12), ('角色', 6), ('測試資料', 24), ('發現', 60), ('類型', 12), ('嚴重度', 8), ('討論結論', 50), ('處置', 10), ('對應版本／commit', 16), ('狀態', 8)]
log = [
    ['v2.6', '2026-10-11', 'B', 'variants/G_xlayer.xlsx', '比較欄被自動判成 Product，Lot_Type（BSL／CIP）沒被選中；戰報變成 PA vs PB，沒有任何 HL。', '欄位抓不到', '高', '自動設定改成先看欄位的「值」（BSL／CIP／POR／ref…）再看欄名；Product 這類名稱不當比較欄的候選。', '已修', 'v2.6 820c2e2', '關閉'],
    ['v2.6', '2026-10-11', 'C', 'variants/G_xlayer.xlsx', 'Rs_M3 的配對圖畫的是同層 M3 THK（r −0.44，對不上），真正解釋偏移的 M2 DISH 沒有圖。', '可更方便', '中', '配對圖優先畫對上的同層 inline；同層都對不上時改畫最強的跨層項目，標「跨層，方向視 pattern 而定」。', '已修', 'v2.6', '關閉'],
    ['v2.6', '2026-10-11', 'C', 'variants/G_xlayer.xlsx', '跨層卡的 r 用 M2 THK（−0.39），圖卻用 M2 DISH（0.57），兩者對不上。', '結論一致性', '中', '卡片與圖都取 |K| 最大的跨層項目。', '已修', 'v2.6', '關閉'],
    ['v2.6', '2026-10-11', 'C', 'WAT_FAB_compare.xlsx', '跨層的建議檢查文字太長，折成兩行、✔ 單獨一行。', '可更方便', '低', '檢查文字縮短為一行（「M2 THK／dishing；WAT 結構與 inline monitor 是否同 pattern」）。', '已修', 'v2.6', '關閉'],
    ['v2.6', '2026-10-11', 'C', 'WAT_FAB_compare.xlsx', 'toast 顯示「182 ms」這種開發資訊，使用者不需要。', '可更方便', '低', '拿掉毫秒。', '已修', 'v2.6', '關閉'],
    ['v2.6', '2026-10-11', 'C', 'WAT_FAB_compare.xlsx', '沒有機台欄時，機台卡顯示「○ 無機台／chamber 欄」，分析照常；符合「缺欄位不擋分析」的要求。', '確認通過', '—', '維持。', '不做', 'v2.6', '關閉'],
    ['v2.7', '2026-10-11', 'C', 'WAT_FAB_compare.xlsx', '「Barrier／liner 偏厚」被關鍵字「偏厚」誤歸到 THK 證據，因 M3 THK 未偏移而被灰掉；barrier 厚度不是 Cu THK monitor 能證明或排除的。', '結論一致性', '中', 'THK 證據的關鍵字拿掉「偏薄／偏厚」；barrier 類原因維持中性（無資料）。', '已修', 'v2.7 d655d38', '關閉'],
    ['v2.7', '2026-10-11', 'A', '知識庫 Excel round-trip 測試', '原因文字裡用了「；」，知識庫 Excel 匯出再匯入時被拆成兩條規則差異。', '卡住', '中', '清單欄的分隔符是「；」或換行，原因文字一律改用「，」；寫進 BEOL_REPORT.md 0.9。', '已修', 'v2.7', '關閉'],
    ['v2.7', '2026-10-11', 'C', 'variants/G_xlayer.xlsx', '動態加入的「M2 topography 繼承」與 base 規則裡的「下層 topography 繼承」重複出現。', '結論一致性', '低', '跨層原因只放在 base 規則，由證據排序把它推到最前，不再動態加。', '已修', 'v2.7', '關閉'],
    ['v2.7', '2026-10-11', 'A', '—', '跨層表 App.XLAYER 目前只能改原始碼，工程師無法在知識庫頁維護。', '分析需求', '中', '下一版：知識庫 Excel 多一個「跨層」sheet，匯入時合併進 XLAYER。', '延後', '待排', '開啟'],
    ['v2.7', '2026-10-11', 'A', '—', 'WAT 結構與 inline monitor 位置不同時只能用文字提醒；若資料有 site／座標，可做「同位置 r vs 不同位置 r」。', '分析需求', '中', '等使用者提供含座標的真實資料格式再設計。', '延後', '待排', '開啟'],
    ['v2.8', '2026-10-12', 'C', '使用者實際資料（截圖）', 'NOVA_THK_R（厚度 range）被自動歸到 inline THK，和均值 THK 一起參與同層配對。', '欄位抓不到', '中', '新增「其他（inline）」「其他（yield）」類別；_R／_SD／_MAX／(R) 等衍生統計後綴自動歸「其他」，仍點燈但不配對；設定頁可手動改類別。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'C', '使用者實際資料（截圖）', '「可能原因」與「建議檢查」兩欄概念重疊、文字量大；需人工與缺資料分兩種圖示但使用者想做的事一樣（補資料看相關）。', '可更方便', '高', '合併成一個「建議檢查」清單：每列狀態（✔ 有關聯／△ 弱關聯／✘ 關聯性不大／○ 缺資料）＋證據＋🔍；🔍 開「檢查視窗」：X／Y 可換、r 與方向判定、共振 trend、「＋ 匯入資料補進」（by wafer id 合併後視窗自動重開）。可能原因降為一行「物理機制」。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'C', '使用者實際資料（截圖）', '鄰層只有 THK／THK range 資料、r = 0.11，卻寫成「下層 topography 繼承 ✔ 資料支持」，過度推論；pattern 資訊系統根本沒有。', '結論一致性', '高', '所有同層／鄰層判定改用 K 與 wafer 間 r 兩者：有偏移且 |r|≥0.3 才是 ✔「有關聯」，否則 △「關聯弱」；文字只寫資料看到的（「鄰層 M14 THK +0.68σ 關聯弱，證據不足」），機制文字只留在物理機制行與 tooltip。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'C', 'variants/H_longnames.xlsx', '「M15 post-CMP THK」被標 ✔ 有關聯（r = −0.58），但 M15 THK 在 CIP 根本沒偏移——相關存在不代表它解釋這次偏移。', '結論一致性', '高', '檢查列的狀態同時要求「該 item 在 CIP 有偏移」與「wafer 間有關聯」；未偏移 → ✘「此 item 未偏移」。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'C', 'variants/H_longnames.xlsx', '機台檢查「已掃 1 個欄位，沒有明顯 bad tool」卻給 ✔。', '結論一致性', '中', '機台類檢查：有 bad tool 才 ✔，掃過沒集中 → ✘。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'C', 'variants/H_longnames.xlsx', '共振 trend 的 CIP 淡底在 BSL／CIP 時間交錯時蓋住整張圖。', '可更方便', '低', 'CIP 區間超過 70% 範圍時不畫淡底，改以 CIP wafer 的點標示。', '已修', 'v2.8', '關閉'],
    ['v2.8', '2026-10-12', 'B', 'variants/H_longnames.xlsx', '長欄名「Mean(Item Value by Site, 07504.050 M15CUCMP_5, NOVA_THK)」整串顯示在卡片與圖軸上。', '可更方便', '中', 'parseName 拿掉統計包裝、固定字樣、站點代碼 → 顯示「M15CUCMP_5 NOVA_THK」；層別仍從 M15CUCMP_5 判出 M15。', '已修', 'v2.8', '關閉'],
    ['v2.9-cm', '2026-10-12', 'A', 'Design／Slides 設計稿', '把 Design 畫布與 Slides 的版面移植成獨立資料夾 BEOL_Analysis_Claude module：theme_cm.css 覆寫外觀（IBM Plex 字體、條件列合併標題／分組／期間／wafer 數／門檻／燈號數、HL 清單兩欄、交叉比對卡兩行式、建議檢查兩欄 tinted rows、視窗 12px 圓角＋callout）；PPT 封面改模組卡、點燈表頁加橘／黃／讀法三張卡。分析、DOM、測試與原版相同。', '分析需求', '—', '原版資料夾不動；兩邊共用引擎，之後只要把 theme_cm.css 與 buildPpt 的封面／點燈表段落搬回去即可合併。', '已修', 'v2.9-cm', '關閉'],
    ['v2.9-cm', '2026-10-12', 'C', 'WAT_FAB_compare.xlsx', '1280 寬時品牌字折成三行、模組切換被截掉；戰報頁上方的大標題與條件列重複。', '可更方便', '中', '工具列可換行、品牌與模組不縮；大標題隱藏，條件列已含標題資訊。', '已修', 'v2.9-cm', '關閉'],
    ['v2.9-cm', '2026-10-12', 'C', 'WAT_FAB_compare.xlsx（手機寬）', '分頁按鈕折成直排；點燈表排在 HL 之前。', '可更方便', '低', '分頁橫向捲動；手機寬時 HL 在上、點燈表在下（同設計稿）。', '已修', 'v2.9-cm', '關閉'],
    ['v2.9-cm', '2026-10-12', 'C', 'WAT_FAB_compare.xlsx', '檢查視窗散佈圖的 Plotly modebar 疊在圖例上。', '可更方便', '低', '關掉 modebar、圖例移到圖上方（兩個版本都修）。', '已修', 'v2.9-cm／v2.9', '關閉'],
    ['v2.9-cm', '2026-10-12', 'B', 'variants/H_longnames.xlsx', '以不同資料形狀（兩組、長欄名、range 欄）跑 CM 版：條件列、卡片、檢查清單、PPT 皆正常。', '確認通過', '—', '維持。', '不做', 'v2.9-cm', '關閉'],
    ['v2.9', '2026-10-12', 'C', '使用者 PPT 截圖（v2.6）', '封面與 HL 標題寫「金屬線變細或變薄」這類意義推論；沒有資料佐證，報告只該描述現象；封面字多。', '結論一致性', '高', '所有標題改成現象：「Con_M15 @ TMQI55：FAB15 +2.30σ」；意義文字改放 dx.mean，只在網頁 HL 清單以灰字顯示。封面 HL 改成每個燈號一行（橘 3：…、…；黃 3：…）。', '已修', 'v2.9', '關閉'],
    ['v2.9', '2026-10-12', 'C', '使用者 PPT 截圖（v2.6）', '點燈表長欄名折三行；「其他」類別排在最前。', '可更方便', '中', '欄名縮寫（v2.8）＋「其他」類別一律排最後；截斷的 Mean(Con_M15 也能縮成 Con_M15。', '已修', 'v2.9', '關閉'],
    ['v2.9', '2026-10-12', 'A', '使用者提議', '12 個 HL 各一頁太多，想一頁多個。', '分析需求', '中', '新增「HL 總表」頁（每列：燈、item@組、偏移、Inline／鄰層／機台／Yield 一句話、先查什麼；8 列一頁）；細節頁預設只做橘燈（HL ≤ 4 時全做），按鈕旁可改「只橘燈／全部」。12 個 HL（橘 4）→ 2 頁總表＋4 頁細節。', '已修', 'v2.9', '關閉'],
    ['v2.9', '2026-10-12', 'C', 'WAT_FAB_compare.xlsx', 'SPA（異常比例）的 r 天生很低（0/1 性質），被判成「關聯弱」。', '結論一致性', '中', 'rate 類 item 不算 r，只用 K 判定。', '已修', 'v2.9', '關閉'],
    ['v2.9', '2026-10-12', 'C', 'variants/H_longnames.xlsx', 'PPT 配對圖註記「r = −0.58，有關聯」但 M15 THK 未偏移。', '結論一致性', '中', '配對圖註記改用同一套 relOf：「此 item 未偏移」。', '已修', 'v2.9', '關閉'],
    ['v2.8', '2026-10-12', 'A', '—', '同層 inline 欄位靠規則自動配對；配錯時使用者只能在檢查視窗換 X，沒有「記住這個對應」。', '分析需求', '中', '下一版：檢查視窗加「設為此 item 的對應欄」，存到 fab 層知識庫（pairs 欄）。', '延後', '待排', '開啟'],
    ['v2.7', '2026-10-11', 'A', '—', '每週 run 貨圖已有時間軸，可加 CUSUM 變化點「偏移從哪一週開始」，對到機台保養／recipe 變更日期。', '分析需求', '低', '列入 P3。', '延後', '待排', '開啟'],
]
ws = sheet(wb, '自審紀錄', log_cols, log, first=False)
dv1 = DataValidation(type='list', formula1='"A,B,C"', allow_blank=True); ws.add_data_validation(dv1); dv1.add('C2:C500')
dv2 = DataValidation(type='list', formula1='"卡住,欄位抓不到,可更方便,結論一致性,分析需求,確認通過"', allow_blank=True); ws.add_data_validation(dv2); dv2.add('F2:F500')
dv3 = DataValidation(type='list', formula1='"高,中,低,—"', allow_blank=True); ws.add_data_validation(dv3); dv3.add('G2:G500')
dv4 = DataValidation(type='list', formula1='"已修,延後,不做"', allow_blank=True); ws.add_data_validation(dv4); dv4.add('I2:I500')
dv5 = DataValidation(type='list', formula1='"開啟,關閉"', allow_blank=True); ws.add_data_validation(dv5); dv5.add('K2:K500')
ws.auto_filter.ref = 'A1:K500'

# ---- 3. 模擬資料
sim = [
    ['v2.8', 'variants/H_longnames.xlsx ＋ H_link_M15CD.xlsx', 'gen_variants.py', '模仿使用者實際資料：長欄名 Mean(Item Value by Site, 站點, NOVA_THK)、NOVA_THK_R range 欄、條件值 FAB18_U%_BSL／FAB15_U%_BSL、M14 CMP tool 欄；Con_M15 +2σ、M15 THK 正常、M14 THK +0.7σ 但 r≈0.1；沒有 CD 欄。第二檔 M15 CD（80% wafer、id 寫法 UAB001-02）供檢查視窗匯入補進。', '使用者截圖', '其他類別、長欄名、鄰層關聯弱措辭、檢查視窗匯入後自動 ✔ 皆正確；抓到 THK ✔ 誤判與機台 ✔ 誤判（已修）。'],
    ['v2.6', 'variants/G_xlayer.xlsx', 'gen_variants.py', 'Rs_M3 偏高（+2.2σ）、Rc_V2 偏高；同層 M3 CD／THK 正常；M2 THK 偏薄、M2 DISH 變大（下層 topography 繼承）；完全沒有機台欄。', 'REFERENCES #11 #12 #13', '比較欄誤判（已修）；跨層卡、配對圖、原因排序皆正確。'],
    ['v2.5', 'TOOL_history_wide / long.xlsx', 'gen_tool_history.py', '機台履歷寬表與長表（含 rework 兩列），wafer id 寫法 UAB123-01 對 UAB123.01；LK_M12 由 ECP02 造成。', '—', '配對率 96%、長表自動轉寬表、K 不因連結而改變。'],
    ['v2.3', 'WAT_FAB_compare.xlsx', 'gen_wat_fab.py', 'TMTQ11 Rc_V12 偏高（V12 CD 偏小，CH_B 造成）、TMRX12 LK／SPA 異常（M12 CD 偏大）、TMWF54 Con_M13 偏高（M13 THK 偏薄，CMP03 與條件混淆）；含 yield／bin。', '—', '主測試資料；81 項 e2e 以它為基準。'],
    ['v2.2', 'variants/F_big_6x6.xlsx', 'gen_big.py', '6 產品 × 6 層 × 2 CIP 的大點燈表。', '—', '自適應熱圖排版。'],
    ['v2.1', 'variants/C／D／E', 'gen_variants.py', 'CSV 與「NAME (unit)」欄名、三個條件；長表＋不同欄名、條件＝Chamber；只有一個條件。', '—', '泛用性：欄名、格式、條件數。'],
]
sheet(wb, '模擬資料', [('版本', 8), ('檔案', 30), ('產生器', 20), ('埋入的故事', 70), ('取材文獻', 20), ('驗證結果', 50)], sim, first=False)

# ---- 4. B／C 檢核清單
chk = [
    ['B', '1', '用新資料直接匯入（Excel／CSV、寬表或長表），不改欄名。', '匯入後自動進到戰報頁，沒有錯誤。'],
    ['B', '2', '確認自動設定：wafer id、日期、比較欄、BSL、CIP、分組欄。', '每一項都是預期的欄位；不對就記「欄位抓不到」。'],
    ['B', '3', '確認 item 類別（Rs／Rc／LK／SPA／CD／THK／topo／yield）。', '分錯類別會影響燈號計算方式（K、倍數、固定值）。'],
    ['C', '4', '看點燈表：異常格是否就是故事埋的那幾格；有沒有不該亮的。', '偽陽性／偽陰性各記一列。'],
    ['C', '5', '點每個 HL：交叉比對卡（Inline／跨層／機台／Yield）是否對上故事。', '卡片文字一行內能看懂。'],
    ['C', '6', '可能原因：✔ 的是否真的有證據；✘ 的理由是否合理；中性的不應被亂判。', '證據排序的核心驗證。'],
    ['C', '7', '建議檢查：每條都能一鍵開到圖或連結檔案；◇ 的是工具外。', '不應有「點了沒反應」。'],
    ['C', '8', '機台分析視窗：結論行、run 貨 %、每週 run 貨、box plot 的 K 標註。', '沒有機台欄時卡片顯示 ○，分析照常。'],
    ['C', '9', '切模組（WAT／Inline／Yield）、換分組欄、換 BSL／CIP，結果立刻重算。', '不應需要重新匯入。'],
    ['C', '10', '匯出 PPT／HTML／複製結論：內容與畫面一致（含原因排序）。', '主管拿到的版本要和畫面一樣。'],
    ['C', '11', '管理者視角：只看 HL 標題＋卡片＋第一條原因，能不能在 30 秒內說出「問題在哪、下一步查什麼」。', '不行就記「可更方便」。'],
]
sheet(wb, '檢核清單', [('角色', 6), ('步驟', 6), ('檢查', 70), ('通過標準', 50)], chk, first=False)

wb.save(OUT)
print('wrote', OUT)
